# Session Feedback Summary Creator
# Edit this file to change this dropdown section.
# app2627.py calls render() on every Streamlit rerun.
# The original section body is preserved below.

from __future__ import annotations
import re
import streamlit as st
import pandas as pd
from io import BytesIO
from docx import Document
import pytz

import re
from dataclasses import dataclass
from io import BytesIO
from typing import Iterable
from urllib.parse import quote_plus


def render():
    st.header("📋 Session Feedback Summary Creator")

    import io
    import re
    import html
    import zipfile
    import hashlib
    import pandas as pd

    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn


    # =========================================================
    # CURRENT REDCap QUESTION SET
    # =========================================================
    # Only these fields will be used, even if an older CSV still
    # contains q005, q006, q007, q008, etc.

    DEFAULT_QUESTION_LABELS = {
        "q001": "The presenter communicated the material clearly and effectively.",
        "q002": "The cases, examples, or clinical reasoning used in this session enhanced my understanding of the material.",
        "q003": "I gained knowledge or skills from this session that I expect to use in future clinical encounters.",
        "q004": "Overall, I found this teaching session valuable to my learning.",
    }

    DEFAULT_COMMENT_LABELS = {
        "session_c001": "What is one aspect of the session that contributed most to your learning?",
        "session_c002": "What suggestions would you offer to further enhance the teaching session?",
    }

    ACTIVE_QUESTION_FIELDS = [
        "q001",
        "q002",
        "q003",
        "q004",
    ]

    ACTIVE_COMMENT_FIELDS = [
        "session_c001",
        "session_c002",
    ]


    # =========================================================
    # GENERAL HELPERS
    # =========================================================

    def get_academic_year(value):
        """
        Academic year begins July 1.
        """

        d = pd.to_datetime(
            value,
            errors="coerce"
        )

        if pd.isna(d):
            return "Unknown"

        if d.month >= 7:
            return f"{d.year}-{d.year + 1}"

        return f"{d.year - 1}-{d.year}"


    def clean_redcap_comment(value):
        """
        Remove REDCap HTML formatting from narrative comments.
        """

        if pd.isna(value):
            return ""

        text = html.unescape(
            str(value)
        )

        text = re.sub(
            r"<br\s*/?>",
            "\n",
            text,
            flags=re.IGNORECASE
        )

        text = re.sub(
            r"</p\s*>",
            "\n",
            text,
            flags=re.IGNORECASE
        )

        text = re.sub(
            r"<[^>]+>",
            "",
            text
        )

        text = text.replace(
            "\xa0",
            " "
        )

        text = re.sub(
            r"[ \t]+",
            " ",
            text
        )

        text = re.sub(
            r"\n\s*\n+",
            "\n",
            text
        )

        return text.strip()


    def safe_filename(value):

        value = re.sub(
            r"[^A-Za-z0-9._ -]",
            "",
            str(value)
        )

        value = value.strip().replace(
            " ",
            "_"
        )

        return value or "Presenter"


    # =========================================================
    # WORD FORMATTING HELPERS
    # =========================================================

    def shade_cell(
        cell,
        fill="D9E1F2"
    ):

        tc_pr = cell._tc.get_or_add_tcPr()

        shading = tc_pr.find(
            qn("w:shd")
        )

        if shading is None:

            shading = OxmlElement(
                "w:shd"
            )

            tc_pr.append(
                shading
            )

        shading.set(
            qn("w:fill"),
            fill
        )


    def set_cell_text(
        cell,
        text,
        bold=False,
        size=9,
        alignment="center"
    ):

        cell.text = ""

        paragraph = cell.paragraphs[0]

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.LEFT
            if alignment == "left"
            else WD_ALIGN_PARAGRAPH.CENTER
        )

        paragraph.paragraph_format.space_after = Pt(
            0
        )

        run = paragraph.add_run(
            str(text)
        )

        run.bold = bold
        run.font.name = "Arial"
        run.font.size = Pt(
            size
        )

        cell.vertical_alignment = (
            WD_CELL_VERTICAL_ALIGNMENT.CENTER
        )


    def add_section_heading(
        document,
        text
    ):

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_before = Pt(
            8
        )

        paragraph.paragraph_format.space_after = Pt(
            3
        )

        run = paragraph.add_run(
            text
        )

        run.bold = True
        run.underline = True
        run.font.name = "Arial"
        run.font.size = Pt(
            11
        )


    # =========================================================
    # PREPARE ONE PRESENTER'S DATA
    # =========================================================

    def prepare_presenter_data(
        evaluation_df,
        presenter_name,
        question_columns
    ):

        presenter_df = evaluation_df[
            evaluation_df[
                "presenter"
            ]
            .astype(str)
            .str.strip()
            == presenter_name
        ].copy()


        # -----------------------------------------------------
        # Convert rating questions to numeric
        # -----------------------------------------------------

        for col in question_columns:

            presenter_df[
                col
            ] = pd.to_numeric(
                presenter_df[
                    col
                ],
                errors="coerce"
            )


        # -----------------------------------------------------
        # Session date
        # -----------------------------------------------------

        presenter_df[
            "_session_date"
        ] = pd.to_datetime(
            presenter_df[
                "date"
            ],
            errors="coerce"
        )


        # -----------------------------------------------------
        # If session date missing, use REDCap timestamp
        # -----------------------------------------------------

        if (
            "form_1_timestamp"
            in presenter_df.columns
        ):

            timestamp_date = pd.to_datetime(
                presenter_df[
                    "form_1_timestamp"
                ],
                errors="coerce"
            )

            presenter_df[
                "_session_date"
            ] = presenter_df[
                "_session_date"
            ].fillna(
                timestamp_date
            )


        # -----------------------------------------------------
        # Academic year
        # -----------------------------------------------------

        presenter_df[
            "_academic_year"
        ] = presenter_df[
            "_session_date"
        ].apply(
            get_academic_year
        )


        # -----------------------------------------------------
        # Clean session title
        # -----------------------------------------------------

        presenter_df[
            "_title"
        ] = (
            presenter_df[
                "title"
            ]
            .fillna(
                "Untitled Session"
            )
            .astype(str)
            .str.strip()
        )

        presenter_df.loc[
            presenter_df[
                "_title"
            ] == "",
            "_title"
        ] = "Untitled Session"


        # -----------------------------------------------------
        # Session identifier
        #
        # Same presenter + same title + same date
        # = one individual teaching session.
        # -----------------------------------------------------

        presenter_df[
            "_session_day"
        ] = presenter_df[
            "_session_date"
        ].dt.strftime(
            "%Y-%m-%d"
        )

        presenter_df[
            "_session_day"
        ] = presenter_df[
            "_session_day"
        ].fillna(
            "Unknown"
        )

        presenter_df[
            "_session_key"
        ] = (
            presenter_df[
                "_title"
            ]
            + "||"
            + presenter_df[
                "_session_day"
            ]
        )

        return presenter_df


    # =========================================================
    # CREATE ONE ROW FOR EACH INDIVIDUAL SESSION
    # =========================================================

    def create_session_summary(
        presenter_df,
        question_columns
    ):

        summary_rows = []


        grouped = presenter_df.groupby(
            [
                "_academic_year",
                "_session_day",
                "_title"
            ],
            dropna=False
        )


        for (
            academic_year,
            session_day,
            session_title
        ), group in grouped:


            # -------------------------------------------------
            # Session rating average
            # -------------------------------------------------

            if question_columns:

                scores = (
                    group[
                        question_columns
                    ]
                    .stack()
                    .dropna()
                )

            else:

                scores = pd.Series(
                    dtype=float
                )


            rating_average = (
                float(
                    scores.mean()
                )
                if len(scores) > 0
                else None
            )


            # -------------------------------------------------
            # Number of evaluations
            # -------------------------------------------------

            number_evaluations = len(
                group
            )


            # -------------------------------------------------
            # Participants
            # -------------------------------------------------

            participant_number = None


            if (
                "par_number"
                in group.columns
            ):

                participant_values = (
                    pd.to_numeric(
                        group[
                            "par_number"
                        ],
                        errors="coerce"
                    )
                    .dropna()
                )


                if len(
                    participant_values
                ) > 0:

                    participant_number = int(
                        participant_values.iloc[
                            0
                        ]
                    )


            # -------------------------------------------------
            # Username
            # -------------------------------------------------

            username = ""


            if (
                "username"
                in group.columns
            ):

                username_values = (
                    group[
                        "username"
                    ]
                    .dropna()
                    .astype(str)
                    .str.strip()
                )


                username_values = [
                    value
                    for value
                    in username_values
                    if value
                ]


                if username_values:

                    username = (
                        username_values[
                            0
                        ]
                    )


            # -------------------------------------------------
            # Friendly date
            # -------------------------------------------------

            parsed_date = pd.to_datetime(
                session_day,
                errors="coerce"
            )


            display_date = (
                ""
                if pd.isna(
                    parsed_date
                )
                else parsed_date.strftime(
                    "%m/%d/%Y"
                )
            )


            summary_rows.append(
                {
                    "Username":
                        username,

                    "Academic Year":
                        academic_year,

                    "Session Date":
                        display_date,

                    "Session Title":
                        session_title,

                    "Participants":
                        participant_number,

                    "Rating Average":
                        rating_average,

                    "Number of Evaluations":
                        number_evaluations,

                    "_session_key":
                        (
                            f"{session_title}"
                            f"||{session_day}"
                        ),
                }
            )


        summary_df = pd.DataFrame(
            summary_rows
        )


        # -----------------------------------------------------
        # Newest session first
        # -----------------------------------------------------

        if not summary_df.empty:

            summary_df[
                "_sort_date"
            ] = pd.to_datetime(
                summary_df[
                    "Session Date"
                ],
                errors="coerce"
            )


            summary_df = (
                summary_df.sort_values(
                    "_sort_date",
                    ascending=False
                )
            )


            summary_df = (
                summary_df.drop(
                    columns=[
                        "_sort_date"
                    ]
                )
            )


        return summary_df


    # =========================================================
    # COLLECT COMMENTS
    # =========================================================

    def collect_presenter_comments(
        presenter_df,
        comment_columns
    ):

        cleaned_comments = {}

        evaluations_with_comments = pd.Series(
            False,
            index=presenter_df.index
        )

        total_comments = 0


        for comment_col in (
            comment_columns
        ):

            entries = []


            for idx, row in (
                presenter_df.iterrows()
            ):

                comment = clean_redcap_comment(
                    row.get(
                        comment_col,
                        ""
                    )
                )


                if not comment:
                    continue


                total_comments += 1

                evaluations_with_comments.loc[
                    idx
                ] = True


                session_date = pd.to_datetime(
                    row.get(
                        "_session_date"
                    ),
                    errors="coerce"
                )


                date_text = (
                    ""
                    if pd.isna(
                        session_date
                    )
                    else session_date.strftime(
                        "%m/%d/%Y"
                    )
                )


                entries.append(
                    {
                        "date":
                            date_text,

                        "title":
                            row.get(
                                "_title",
                                ""
                            ),

                        "comment":
                            comment,
                    }
                )


            cleaned_comments[
                comment_col
            ] = entries


        return (
            cleaned_comments,
            int(
                evaluations_with_comments.sum()
            ),
            total_comments,
        )


    # =========================================================
    # CREATE ONE PRESENTER'S WORD REPORT
    # =========================================================

    def create_presenter_word_report(
        evaluation_df,
        presenter_name,
        question_columns,
        question_labels,
        comment_columns,
        comment_labels,
        custom_comment_summary=""
    ):

        presenter_df = prepare_presenter_data(
            evaluation_df,
            presenter_name,
            question_columns
        )


        session_summary = create_session_summary(
            presenter_df,
            question_columns
        )


        # =====================================================
        # OVERALL PRESENTER RATING
        # =====================================================

        if question_columns:

            all_scores = (
                presenter_df[
                    question_columns
                ]
                .stack()
                .dropna()
            )

        else:

            all_scores = pd.Series(
                dtype=float
            )


        if len(
            all_scores
        ) > 0:

            overall_rating = float(
                all_scores.mean()
            )

            overall_rating_text = (
                f"{overall_rating:.2f}/5.00"
            )

        else:

            overall_rating_text = "N/A"


        total_evaluations = len(
            presenter_df
        )


        (
            cleaned_comments,
            number_evaluations_with_comments,
            total_comments,
        ) = collect_presenter_comments(
            presenter_df,
            comment_columns
        )


        # =====================================================
        # CREATE WORD DOCUMENT
        # =====================================================

        document = Document()


        section = document.sections[
            0
        ]

        section.top_margin = Inches(
            0.55
        )

        section.bottom_margin = Inches(
            0.55
        )

        section.left_margin = Inches(
            0.55
        )

        section.right_margin = Inches(
            0.55
        )


        normal_style = document.styles[
            "Normal"
        ]

        normal_style.font.name = "Arial"

        normal_style.font.size = Pt(
            9
        )


        # =====================================================
        # PRESENTER NAME
        # =====================================================

        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        paragraph.paragraph_format.space_after = Pt(
            2
        )

        run = paragraph.add_run(
            presenter_name
        )

        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(
            16
        )


        # =====================================================
        # DOCUMENT TITLE
        # =====================================================

        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        paragraph.paragraph_format.space_after = Pt(
            7
        )

        run = paragraph.add_run(
            "Teaching Evaluation Summary"
        )

        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(
            12
        )


        # =====================================================
        # OVERALL SUMMARY
        # =====================================================

        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        paragraph.paragraph_format.space_after = Pt(
            8
        )

        run = paragraph.add_run(
            f"Completed Evaluations: "
            f"{total_evaluations}"
            f"    |    "
            f"Overall Rating: "
            f"{overall_rating_text}"
        )

        run.font.name = "Arial"

        run.font.size = Pt(
            9
        )


        # =====================================================
        # TEACHING ACTIVITY SUMMARY
        # =====================================================

        add_section_heading(
            document,
            "Teaching Activity Summary"
        )


        activity_table = document.add_table(
            rows=1,
            cols=6
        )

        activity_table.style = (
            "Table Grid"
        )

        activity_table.alignment = (
            WD_TABLE_ALIGNMENT.CENTER
        )

        activity_table.autofit = True


        activity_headers = [
            "Academic Year",
            "Session Date",
            "Session Title",
            "Participants",
            "Rating Average\nLow 1 - High 5",
            "Number of\nEvaluations",
        ]


        for index, header in enumerate(
            activity_headers
        ):

            shade_cell(
                activity_table.rows[
                    0
                ].cells[
                    index
                ]
            )


            set_cell_text(
                activity_table.rows[
                    0
                ].cells[
                    index
                ],
                header,
                bold=True,
                size=8
            )


        for _, row in (
            session_summary.iterrows()
        ):

            cells = (
                activity_table
                .add_row()
                .cells
            )


            rating = row[
                "Rating Average"
            ]


            rating_text = (
                "N/A"
                if pd.isna(
                    rating
                )
                else f"{rating:.2f}"
            )


            participants = row[
                "Participants"
            ]


            participant_text = (
                ""
                if pd.isna(
                    participants
                )
                else str(
                    int(
                        participants
                    )
                )
            )


            values = [
                row[
                    "Academic Year"
                ],

                row[
                    "Session Date"
                ],

                row[
                    "Session Title"
                ],

                participant_text,

                rating_text,

                int(
                    row[
                        "Number of Evaluations"
                    ]
                ),
            ]


            for index, value in enumerate(
                values
            ):

                set_cell_text(
                    cells[
                        index
                    ],
                    value,
                    size=8,
                    alignment=(
                        "left"
                        if index == 2
                        else "center"
                    )
                )


        # =====================================================
        # QUESTION AVERAGES — ONLY q001-q004
        # =====================================================

        add_section_heading(
            document,
            "Teaching Evaluations"
        )


        if question_columns:

            question_table = document.add_table(
                rows=1,
                cols=3
            )

            question_table.style = (
                "Table Grid"
            )

            question_table.alignment = (
                WD_TABLE_ALIGNMENT.CENTER
            )


            question_headers = [
                "Evaluation Item",
                (
                    "Average Response\n"
                    "1 = Strongly Disagree\n"
                    "5 = Strongly Agree"
                ),
                "N",
            ]


            for index, header in enumerate(
                question_headers
            ):

                shade_cell(
                    question_table.rows[
                        0
                    ].cells[
                        index
                    ]
                )


                set_cell_text(
                    question_table.rows[
                        0
                    ].cells[
                        index
                    ],
                    header,
                    bold=True,
                    size=8
                )


            for question_col in (
                question_columns
            ):

                values = (
                    pd.to_numeric(
                        presenter_df[
                            question_col
                        ],
                        errors="coerce"
                    )
                    .dropna()
                )


                average_text = (
                    f"{values.mean():.2f}"
                    if len(
                        values
                    ) > 0
                    else "N/A"
                )


                cells = (
                    question_table
                    .add_row()
                    .cells
                )


                set_cell_text(
                    cells[
                        0
                    ],
                    question_labels.get(
                        question_col,
                        question_col
                    ),
                    size=8,
                    alignment="left"
                )


                set_cell_text(
                    cells[
                        1
                    ],
                    average_text,
                    size=8
                )


                set_cell_text(
                    cells[
                        2
                    ],
                    len(
                        values
                    ),
                    size=8
                )


        else:

            paragraph = document.add_paragraph(
                "None of the active rating questions "
                "(q001-q004) were detected."
            )

            paragraph.runs[
                0
            ].font.size = Pt(
                8.5
            )


        # =====================================================
        # NARRATIVE FEEDBACK SUMMARY
        # =====================================================

        add_section_heading(
            document,
            "Narrative Feedback Summary"
        )


        if (
            custom_comment_summary.strip()
        ):

            narrative_summary = (
                custom_comment_summary.strip()
            )

        else:

            narrative_summary = (
                f"Written feedback was provided on "
                f"{number_evaluations_with_comments} "
                f"of {total_evaluations} completed "
                f"evaluations. A total of "
                f"{total_comments} narrative "
                f"comment(s) were submitted."
            )


        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(
            5
        )

        run = paragraph.add_run(
            narrative_summary
        )

        run.font.name = "Arial"

        run.font.size = Pt(
            9
        )


        # =====================================================
        # COMMENT COUNTS
        # =====================================================

        if comment_columns:

            comment_summary_table = (
                document.add_table(
                    rows=1,
                    cols=2
                )
            )

            comment_summary_table.style = (
                "Table Grid"
            )

            comment_summary_table.alignment = (
                WD_TABLE_ALIGNMENT.CENTER
            )


            shade_cell(
                comment_summary_table.rows[
                    0
                ].cells[
                    0
                ]
            )

            shade_cell(
                comment_summary_table.rows[
                    0
                ].cells[
                    1
                ]
            )


            set_cell_text(
                comment_summary_table.rows[
                    0
                ].cells[
                    0
                ],
                "Narrative Feedback Item",
                bold=True,
                size=8
            )


            set_cell_text(
                comment_summary_table.rows[
                    0
                ].cells[
                    1
                ],
                "Comments Submitted",
                bold=True,
                size=8
            )


            for comment_col in (
                comment_columns
            ):

                cells = (
                    comment_summary_table
                    .add_row()
                    .cells
                )


                set_cell_text(
                    cells[
                        0
                    ],
                    comment_labels.get(
                        comment_col,
                        comment_col
                    ),
                    size=8,
                    alignment="left"
                )


                set_cell_text(
                    cells[
                        1
                    ],
                    len(
                        cleaned_comments.get(
                            comment_col,
                            []
                        )
                    ),
                    size=8
                )


        # =====================================================
        # VERBATIM LEARNER COMMENTS
        # =====================================================

        add_section_heading(
            document,
            "Learner Comments"
        )


        comments_written = False


        for comment_col in (
            comment_columns
        ):

            entries = cleaned_comments.get(
                comment_col,
                []
            )


            if not entries:
                continue


            comments_written = True


            paragraph = document.add_paragraph()

            paragraph.paragraph_format.space_before = Pt(
                4
            )

            paragraph.paragraph_format.space_after = Pt(
                1
            )


            run = paragraph.add_run(
                comment_labels.get(
                    comment_col,
                    comment_col
                )
            )

            run.bold = True
            run.font.name = "Arial"

            run.font.size = Pt(
                9
            )


            for entry in entries:

                context = []


                if entry[
                    "date"
                ]:

                    context.append(
                        entry[
                            "date"
                        ]
                    )


                if entry[
                    "title"
                ]:

                    context.append(
                        entry[
                            "title"
                        ]
                    )


                prefix = (
                    " — ".join(
                        context
                    )
                    + ": "
                    if context
                    else ""
                )


                paragraph = document.add_paragraph(
                    style="List Bullet"
                )

                paragraph.paragraph_format.space_after = Pt(
                    1
                )


                run = paragraph.add_run(
                    prefix
                    + entry[
                        "comment"
                    ]
                )

                run.font.name = "Arial"

                run.font.size = Pt(
                    8.5
                )


        if not comments_written:

            paragraph = document.add_paragraph(
                "No narrative comments were submitted."
            )

            paragraph.runs[
                0
            ].font.size = Pt(
                8.5
            )


        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_before = Pt(
            6
        )

        run = paragraph.add_run(
            "Learner comments are presented verbatim "
            "except for removal of REDCap HTML formatting."
        )

        run.italic = True
        run.font.name = "Arial"

        run.font.size = Pt(
            7.5
        )


        # =====================================================
        # SAVE WORD FILE
        # =====================================================

        output = io.BytesIO()

        document.save(
            output
        )

        output.seek(
            0
        )

        return output.getvalue()


    # =========================================================
    # CREATE MASTER CSV FOR ZIP
    # =========================================================
    #
    # One row = one individual teaching session.
    #
    # username is intentionally the FIRST column.
    # =========================================================

    def create_master_summary_csv(
        evaluation_df,
        presenters,
        question_columns,
        question_labels,
        comment_columns,
        comment_labels,
        custom_comment_summaries
    ):

        master_rows = []


        for presenter_name in (
            presenters
        ):

            presenter_df = prepare_presenter_data(
                evaluation_df,
                presenter_name,
                question_columns
            )


            session_summary = create_session_summary(
                presenter_df,
                question_columns
            )


            # =================================================
            # OVERALL PRESENTER RATING
            # =================================================

            if question_columns:

                all_scores = (
                    presenter_df[
                        question_columns
                    ]
                    .stack()
                    .dropna()
                )

            else:

                all_scores = pd.Series(
                    dtype=float
                )


            overall_rating = (
                round(
                    float(
                        all_scores.mean()
                    ),
                    2
                )
                if len(
                    all_scores
                ) > 0
                else ""
            )


            total_evaluations = len(
                presenter_df
            )


            # =================================================
            # PRESENTER-WIDE QUESTION AVERAGES
            # =================================================

            question_results = {}


            for question_col in (
                question_columns
            ):

                values = (
                    pd.to_numeric(
                        presenter_df[
                            question_col
                        ],
                        errors="coerce"
                    )
                    .dropna()
                )


                question_results[
                    question_col
                ] = {

                    "question":
                        question_labels.get(
                            question_col,
                            question_col
                        ),

                    "average":
                        (
                            round(
                                float(
                                    values.mean()
                                ),
                                2
                            )
                            if len(
                                values
                            ) > 0
                            else ""
                        ),

                    "n":
                        len(
                            values
                        ),
                }


            # =================================================
            # PRESENTER-WIDE COMMENTS
            # =================================================

            evaluations_with_comments = pd.Series(
                False,
                index=presenter_df.index
            )

            total_comments = 0

            presenter_comment_counts = {}


            for comment_col in (
                comment_columns
            ):

                this_prompt_count = 0


                for idx, row in (
                    presenter_df.iterrows()
                ):

                    comment = clean_redcap_comment(
                        row.get(
                            comment_col,
                            ""
                        )
                    )


                    if not comment:
                        continue


                    total_comments += 1
                    this_prompt_count += 1

                    evaluations_with_comments.loc[
                        idx
                    ] = True


                presenter_comment_counts[
                    comment_col
                ] = this_prompt_count


            number_evaluations_with_comments = int(
                evaluations_with_comments.sum()
            )


            # =================================================
            # NARRATIVE SUMMARY
            # =================================================

            custom_summary = (
                custom_comment_summaries.get(
                    presenter_name,
                    ""
                )
                .strip()
            )


            if custom_summary:

                narrative_summary = (
                    custom_summary
                )

            else:

                narrative_summary = (
                    f"Written feedback was provided on "
                    f"{number_evaluations_with_comments} "
                    f"of {total_evaluations} completed "
                    f"evaluations. A total of "
                    f"{total_comments} narrative "
                    f"comment(s) were submitted."
                )


            # =================================================
            # ONE CSV ROW PER INDIVIDUAL SESSION
            # =================================================

            for _, session_row in (
                session_summary.iterrows()
            ):

                output_row = {}


                # ---------------------------------------------
                # USERNAME MUST BE FIRST
                # ---------------------------------------------

                output_row[
                    "username"
                ] = session_row[
                    "Username"
                ]


                # ---------------------------------------------
                # Session information
                # ---------------------------------------------

                output_row[
                    "presenter"
                ] = presenter_name


                output_row[
                    "academic_year"
                ] = session_row[
                    "Academic Year"
                ]


                output_row[
                    "session_date"
                ] = session_row[
                    "Session Date"
                ]


                output_row[
                    "session_title"
                ] = session_row[
                    "Session Title"
                ]


                participants = session_row[
                    "Participants"
                ]


                output_row[
                    "participants"
                ] = (
                    ""
                    if pd.isna(
                        participants
                    )
                    else int(
                        participants
                    )
                )


                session_rating = session_row[
                    "Rating Average"
                ]


                output_row[
                    "session_rating_average"
                ] = (
                    ""
                    if pd.isna(
                        session_rating
                    )
                    else round(
                        float(
                            session_rating
                        ),
                        2
                    )
                )


                output_row[
                    "number_of_evaluations"
                ] = int(
                    session_row[
                        "Number of Evaluations"
                    ]
                )


                output_row[
                    "overall_presenter_rating"
                ] = overall_rating


                output_row[
                    "total_presenter_evaluations"
                ] = total_evaluations


                output_row[
                    "rating_scale"
                ] = (
                    "1 = Strongly Disagree; "
                    "5 = Strongly Agree"
                )


                # ---------------------------------------------
                # ONLY q001-q004
                # ---------------------------------------------

                for question_col in (
                    question_columns
                ):

                    result = question_results[
                        question_col
                    ]


                    output_row[
                        f"{question_col}_question"
                    ] = result[
                        "question"
                    ]


                    output_row[
                        f"{question_col}_average"
                    ] = result[
                        "average"
                    ]


                    output_row[
                        f"{question_col}_n"
                    ] = result[
                        "n"
                    ]


                # ---------------------------------------------
                # Narrative summary
                # ---------------------------------------------

                output_row[
                    "narrative_feedback_summary"
                ] = narrative_summary


                output_row[
                    "evaluations_with_comments"
                ] = (
                    number_evaluations_with_comments
                )


                output_row[
                    "total_comments"
                ] = total_comments


                # ---------------------------------------------
                # Comments from THIS individual session
                # ---------------------------------------------

                session_key = session_row[
                    "_session_key"
                ]


                session_evaluations = presenter_df[
                    presenter_df[
                        "_session_key"
                    ] == session_key
                ]


                for comment_col in (
                    comment_columns
                ):

                    output_row[
                        f"{comment_col}_prompt"
                    ] = comment_labels.get(
                        comment_col,
                        comment_col
                    )


                    output_row[
                        f"{comment_col}_total_comments"
                    ] = (
                        presenter_comment_counts.get(
                            comment_col,
                            0
                        )
                    )


                    session_comments = []


                    for _, evaluation_row in (
                        session_evaluations.iterrows()
                    ):

                        comment = clean_redcap_comment(
                            evaluation_row.get(
                                comment_col,
                                ""
                            )
                        )


                        if comment:

                            session_comments.append(
                                comment
                            )


                    output_row[
                        f"{comment_col}_session_comments"
                    ] = " | ".join(
                        session_comments
                    )


                master_rows.append(
                    output_row
                )


        master_df = pd.DataFrame(
            master_rows
        )


        csv_buffer = io.StringIO()


        master_df.to_csv(
            csv_buffer,
            index=False
        )


        # UTF-8-SIG opens cleanly in Excel
        return csv_buffer.getvalue().encode(
            "utf-8-sig"
        )


    # =========================================================
    # USER INTERFACE
    # =========================================================

    st.write(
        "Upload the REDCap CSV export to create "
        "promotion-ready teaching evaluation summaries "
        "for each presenter."
    )


    st.caption(
        "Each unique session date + session title is "
        "listed as an individual teaching session. "
        "Repeated session titles on different dates "
        "remain separate."
    )


    # =========================================================
    # UPLOAD REDCap CSV
    # =========================================================

    uploaded_file = st.file_uploader(
        "Upload REDCap Evaluation CSV",
        type=[
            "csv"
        ],
        key="session_feedback_summary_upload"
    )


    if uploaded_file is not None:

        file_bytes = (
            uploaded_file.getvalue()
        )


        current_file_hash = hashlib.sha256(
            file_bytes
        ).hexdigest()


        # -----------------------------------------------------
        # Clear previous reports if different CSV uploaded
        # -----------------------------------------------------

        if (
            st.session_state.get(
                "feedback_summary_file_hash"
            )
            != current_file_hash
        ):

            st.session_state[
                "feedback_summary_file_hash"
            ] = current_file_hash


            st.session_state.pop(
                "feedback_summary_reports",
                None
            )


            st.session_state.pop(
                "feedback_summary_zip",
                None
            )


        # =====================================================
        # READ CSV
        # =====================================================

        try:

            evaluation_df = pd.read_csv(
                io.BytesIO(
                    file_bytes
                )
            )


        except Exception as e:

            st.error(
                f"Could not read the REDCap CSV: {e}"
            )

            evaluation_df = None


        if evaluation_df is not None:


            # =================================================
            # REQUIRED COLUMNS
            # =================================================

            required_columns = {
                "username",
                "presenter",
                "title",
                "date",
                "par_number",
            }


            missing_columns = (
                required_columns
                - set(
                    evaluation_df.columns
                )
            )


            if missing_columns:

                st.error(
                    "The REDCap export is missing "
                    "required column(s): "
                    + ", ".join(
                        sorted(
                            missing_columns
                        )
                    )
                )


            else:


                # =============================================
                # COMPLETED EVALUATIONS ONLY
                # =============================================

                if (
                    "form_1_complete"
                    in evaluation_df.columns
                ):

                    evaluation_df[
                        "form_1_complete"
                    ] = pd.to_numeric(
                        evaluation_df[
                            "form_1_complete"
                        ],
                        errors="coerce"
                    )


                    original_count = len(
                        evaluation_df
                    )


                    evaluation_df = (
                        evaluation_df[
                            evaluation_df[
                                "form_1_complete"
                            ] == 2
                        ]
                        .copy()
                    )


                    removed_count = (
                        original_count
                        - len(
                            evaluation_df
                        )
                    )


                    if removed_count > 0:

                        st.caption(
                            f"{removed_count} incomplete or "
                            "unverified evaluation(s) excluded."
                        )


                # =============================================
                # CLEAN PRESENTER NAMES
                # =============================================

                evaluation_df = (
                    evaluation_df[
                        evaluation_df[
                            "presenter"
                        ].notna()
                    ]
                    .copy()
                )


                evaluation_df[
                    "presenter"
                ] = (
                    evaluation_df[
                        "presenter"
                    ]
                    .astype(str)
                    .str.strip()
                )


                evaluation_df = (
                    evaluation_df[
                        evaluation_df[
                            "presenter"
                        ] != ""
                    ]
                    .copy()
                )


                # =============================================
                # ONLY USE q001-q004
                # =============================================
                #
                # q005-q008 are ignored even if present in
                # an older REDCap export.
                # =============================================

                detected_questions = [
                    field
                    for field in ACTIVE_QUESTION_FIELDS
                    if field in evaluation_df.columns
                ]


                # =============================================
                # ONLY USE session_c001 / session_c002
                # =============================================

                detected_comments = [
                    field
                    for field in ACTIVE_COMMENT_FIELDS
                    if field in evaluation_df.columns
                ]


                # =============================================
                # DEFAULT QUESTION / COMMENT MAPPINGS
                # =============================================

                question_columns = (
                    detected_questions.copy()
                )


                question_labels = {
                    field:
                        DEFAULT_QUESTION_LABELS[
                            field
                        ]
                    for field
                    in detected_questions
                }


                comment_columns = (
                    detected_comments.copy()
                )


                comment_labels = {
                    field:
                        DEFAULT_COMMENT_LABELS[
                            field
                        ]
                    for field
                    in detected_comments
                }


                # =============================================
                # WARN IF AN EXPECTED CURRENT FIELD IS MISSING
                # =============================================

                missing_active_questions = [
                    field
                    for field in ACTIVE_QUESTION_FIELDS
                    if field not in evaluation_df.columns
                ]


                if missing_active_questions:

                    st.warning(
                        "The uploaded CSV is missing "
                        "active rating field(s): "
                        + ", ".join(
                            missing_active_questions
                        )
                    )


                missing_active_comments = [
                    field
                    for field in ACTIVE_COMMENT_FIELDS
                    if field not in evaluation_df.columns
                ]


                if missing_active_comments:

                    st.warning(
                        "The uploaded CSV is missing "
                        "active comment field(s): "
                        + ", ".join(
                            missing_active_comments
                        )
                    )


                # =============================================
                # UPLOAD SUMMARY
                # =============================================

                presenter_list = sorted(
                    evaluation_df[
                        "presenter"
                    ]
                    .dropna()
                    .unique()
                    .tolist()
                )


                col1, col2, col3 = st.columns(
                    3
                )


                col1.metric(
                    "Presenters",
                    len(
                        presenter_list
                    )
                )


                col2.metric(
                    "Completed Evaluations",
                    len(
                        evaluation_df
                    )
                )


                col3.metric(
                    "Rating Questions",
                    len(
                        detected_questions
                    )
                )


                # =============================================
                # TROUBLESHOOTING
                # =============================================
                #
                # Mapping tables remain hidden unless needed.
                # =============================================

                with st.expander(
                    "🛠 Troubleshooting",
                    expanded=False
                ):

                    show_question_mapping = st.toggle(
                        "Show REDCap question mapping",
                        value=False,
                        key=(
                            "feedback_summary_"
                            "show_mapping"
                        )
                    )


                    if show_question_mapping:

                        st.caption(
                            "Use this only to verify or "
                            "temporarily adjust how the active "
                            "REDCap fields are mapped to the report."
                        )


                        # =====================================
                        # RATING QUESTIONS — q001-q004 ONLY
                        # =====================================

                        st.markdown(
                            "#### Rating Questions"
                        )


                        if detected_questions:

                            question_config = pd.DataFrame(
                                {
                                    "Include":
                                        [
                                            True
                                            for _ in
                                            detected_questions
                                        ],

                                    "REDCap Field":
                                        detected_questions,

                                    (
                                        "Question / "
                                        "Evaluation Item"
                                    ):
                                        [
                                            DEFAULT_QUESTION_LABELS[
                                                field
                                            ]
                                            for field in
                                            detected_questions
                                        ],
                                }
                            )


                            edited_question_config = (
                                st.data_editor(
                                    question_config,
                                    hide_index=True,
                                    use_container_width=True,
                                    disabled=[
                                        "REDCap Field"
                                    ],
                                    key=(
                                        "feedback_summary_"
                                        "question_editor"
                                    )
                                )
                            )


                            included_question_rows = (
                                edited_question_config[
                                    edited_question_config[
                                        "Include"
                                    ] == True
                                ]
                            )


                            question_columns = (
                                included_question_rows[
                                    "REDCap Field"
                                ]
                                .tolist()
                            )


                            question_labels = dict(
                                zip(
                                    included_question_rows[
                                        "REDCap Field"
                                    ],

                                    included_question_rows[
                                        (
                                            "Question / "
                                            "Evaluation Item"
                                        )
                                    ]
                                )
                            )


                        else:

                            st.warning(
                                "None of q001, q002, "
                                "q003, or q004 were detected."
                            )


                        # =====================================
                        # NARRATIVE COMMENTS
                        # =====================================

                        st.markdown(
                            "#### Narrative Comment Questions"
                        )


                        if detected_comments:

                            comment_config = pd.DataFrame(
                                {
                                    "Include":
                                        [
                                            True
                                            for _ in
                                            detected_comments
                                        ],

                                    "REDCap Field":
                                        detected_comments,

                                    "Comment Prompt":
                                        [
                                            DEFAULT_COMMENT_LABELS[
                                                field
                                            ]
                                            for field in
                                            detected_comments
                                        ],
                                }
                            )


                            edited_comment_config = (
                                st.data_editor(
                                    comment_config,
                                    hide_index=True,
                                    use_container_width=True,
                                    disabled=[
                                        "REDCap Field"
                                    ],
                                    key=(
                                        "feedback_summary_"
                                        "comment_editor"
                                    )
                                )
                            )


                            included_comment_rows = (
                                edited_comment_config[
                                    edited_comment_config[
                                        "Include"
                                    ] == True
                                ]
                            )


                            comment_columns = (
                                included_comment_rows[
                                    "REDCap Field"
                                ]
                                .tolist()
                            )


                            comment_labels = dict(
                                zip(
                                    included_comment_rows[
                                        "REDCap Field"
                                    ],

                                    included_comment_rows[
                                        "Comment Prompt"
                                    ]
                                )
                            )


                        else:

                            st.info(
                                "Neither session_c001 nor "
                                "session_c002 was detected."
                            )


                # =============================================
                # SELECT PRESENTERS
                # =============================================

                st.subheader(
                    "Presenters"
                )


                presenters_to_generate = (
                    st.multiselect(
                        "Presenters to include",
                        options=presenter_list,
                        default=presenter_list,
                        key=(
                            "feedback_summary_"
                            "presenter_selection"
                        )
                    )
                )


                # =============================================
                # OPTIONAL NARRATIVE SUMMARIES
                # =============================================

                custom_comment_summaries = {}


                if presenters_to_generate:

                    with st.expander(
                        "✏️ Optional Narrative Comment Summaries"
                    ):

                        st.caption(
                            "Optional. Leave blank and the "
                            "report will automatically state "
                            "how many evaluations contained "
                            "written feedback and how many "
                            "comments were submitted."
                        )


                        for person in (
                            presenters_to_generate
                        ):

                            custom_comment_summaries[
                                person
                            ] = st.text_area(
                                person,
                                value="",
                                placeholder=(
                                    "Optional promotion-ready "
                                    "summary of recurring themes..."
                                ),
                                key=(
                                    "feedback_comment_summary_"
                                    + safe_filename(
                                        person
                                    )
                                )
                            )


                # =============================================
                # GENERATE REPORTS
                # =============================================

                if st.button(
                    "📄 Create Presenter Word Summaries",
                    type="primary",
                    use_container_width=True,
                    key=(
                        "create_session_feedback_"
                        "summary_reports"
                    )
                ):


                    if not presenters_to_generate:

                        st.warning(
                            "Select at least one presenter."
                        )


                    else:

                        generated_reports = {}


                        # =====================================
                        # CREATE WORD DOCUMENTS
                        # =====================================

                        for presenter_name in (
                            presenters_to_generate
                        ):

                            report_bytes = (
                                create_presenter_word_report(
                                    evaluation_df=(
                                        evaluation_df
                                    ),

                                    presenter_name=(
                                        presenter_name
                                    ),

                                    question_columns=(
                                        question_columns
                                    ),

                                    question_labels=(
                                        question_labels
                                    ),

                                    comment_columns=(
                                        comment_columns
                                    ),

                                    comment_labels=(
                                        comment_labels
                                    ),

                                    custom_comment_summary=(
                                        custom_comment_summaries
                                        .get(
                                            presenter_name,
                                            ""
                                        )
                                    ),
                                )
                            )


                            filename = (
                                safe_filename(
                                    presenter_name
                                )
                                + "_Teaching_Evaluation_"
                                + "Summary.docx"
                            )


                            generated_reports[
                                presenter_name
                            ] = {
                                "filename":
                                    filename,

                                "bytes":
                                    report_bytes,
                            }


                        # =====================================
                        # CREATE MASTER CSV
                        # =====================================

                        master_csv_bytes = (
                            create_master_summary_csv(
                                evaluation_df=(
                                    evaluation_df
                                ),

                                presenters=(
                                    presenters_to_generate
                                ),

                                question_columns=(
                                    question_columns
                                ),

                                question_labels=(
                                    question_labels
                                ),

                                comment_columns=(
                                    comment_columns
                                ),

                                comment_labels=(
                                    comment_labels
                                ),

                                custom_comment_summaries=(
                                    custom_comment_summaries
                                ),
                            )
                        )


                        # =====================================
                        # CREATE ZIP PACKAGE
                        # =====================================

                        zip_buffer = io.BytesIO()


                        with zipfile.ZipFile(
                            zip_buffer,
                            mode="w",
                            compression=(
                                zipfile.ZIP_DEFLATED
                            )
                        ) as zip_file:


                            # ---------------------------------
                            # Add Word reports
                            # ---------------------------------

                            for report_data in (
                                generated_reports.values()
                            ):

                                zip_file.writestr(
                                    report_data[
                                        "filename"
                                    ],

                                    report_data[
                                        "bytes"
                                    ]
                                )


                            # ---------------------------------
                            # Add master CSV
                            # ---------------------------------

                            zip_file.writestr(
                                (
                                    "Presenter_Teaching_"
                                    "Evaluation_Summary.csv"
                                ),

                                master_csv_bytes
                            )


                        zip_buffer.seek(
                            0
                        )


                        # =====================================
                        # STORE FILES IN SESSION STATE
                        # =====================================

                        st.session_state[
                            "feedback_summary_reports"
                        ] = generated_reports


                        st.session_state[
                            "feedback_summary_zip"
                        ] = zip_buffer.getvalue()


                # =============================================
                # DOWNLOAD GENERATED FILES
                # =============================================

                if (
                    "feedback_summary_reports"
                    in st.session_state
                ):

                    reports = st.session_state[
                        "feedback_summary_reports"
                    ]


                    st.success(
                        f"Created {len(reports)} "
                        "presenter Word summary file(s)."
                    )


                    # =========================================
                    # INDIVIDUAL WORD DOWNLOADS
                    # =========================================

                    for (
                        presenter_name,
                        report_data
                    ) in reports.items():


                        st.download_button(
                            label=(
                                "⬇️ Download "
                                f"{presenter_name}"
                            ),

                            data=(
                                report_data[
                                    "bytes"
                                ]
                            ),

                            file_name=(
                                report_data[
                                    "filename"
                                ]
                            ),

                            mime=(
                                "application/"
                                "vnd.openxmlformats-officedocument."
                                "wordprocessingml.document"
                            ),

                            use_container_width=True,

                            key=(
                                "download_feedback_summary_"
                                + safe_filename(
                                    presenter_name
                                )
                            )
                        )


                    # =========================================
                    # COMPLETE ZIP DOWNLOAD
                    # =========================================

                    st.download_button(
                        "📦 Download Complete Summary Package",

                        data=(
                            st.session_state[
                                "feedback_summary_zip"
                            ]
                        ),

                        file_name=(
                            "Presenter_Teaching_"
                            "Evaluation_Summaries.zip"
                        ),

                        mime="application/zip",

                        use_container_width=True,

                        key=(
                            "download_all_"
                            "feedback_summaries"
                        )
                    )
